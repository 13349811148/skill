"""Date replacement in persistent product/monthly store workbooks.

Read existing XLSX values with openpyxl; author and recalculate only with artifact-tool.
"""
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from pathlib import Path
import argparse
import json
import os
import re
import subprocess
import tempfile
import csv

SCRIPTS = Path(__file__).resolve().parent


def day(value):
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d')
    return date.fromisoformat(str(value).strip()[:10]).isoformat()


def requested_dates(arguments):
    dates = set()
    start = end = None
    for i, argument in enumerate(arguments):
        option, equals, value = argument.partition('=')
        if option not in ('--date', '--start-date', '--end-date'):
            continue
        value = value if equals else arguments[i + 1]
        if option == '--date':
            dates.update(day(v) for v in value.split(','))
        elif option == '--start-date':
            start = date.fromisoformat(value)
        else:
            end = date.fromisoformat(value)
    if bool(start) != bool(end) or (start and end < start):
        raise ValueError('开始和结束日期必须成对且顺序正确')
    if start:
        dates.update((start + timedelta(days=i)).isoformat() for i in range((end-start).days+1))
    if not dates:
        raise ValueError('持久报表必须明确请求日期，不能按本次非空行猜覆盖范围')
    return dates


@contextmanager
def locked(output):
    output.parent.mkdir(parents=True, exist_ok=True)
    lock = output.with_suffix(output.suffix + '.lock')
    try:
        descriptor = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError as exc:
        raise ValueError(f'报表正在更新；若先前进程异常退出，请核对后处理锁文件: {lock}') from exc
    os.close(descriptor)
    try:
        yield
    finally:
        lock.unlink(missing_ok=True)


def read_product(output, columns):
    if not output.exists():
        return {}
    from openpyxl import load_workbook
    wb = load_workbook(output, read_only=True, data_only=False)
    months = {}
    try:
        for sheet in wb:
            if not re.fullmatch(r'\d{4}-\d{2}', sheet.title):
                raise ValueError(f'商品总表含旧格式/未知子表 {sheet.title}，请先明确迁移，不覆盖')
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(values_only=True)
            if [v if v is not None else '' for v in next(iterator, ())] != columns:
                raise ValueError(f'商品总表表头不一致: {sheet.title}')
            rows = []
            for raw in iterator:
                if not any(v is not None for v in raw):
                    continue
                if any(isinstance(v, str) and v.startswith('=') for v in raw):
                    raise ValueError('商品总表存在人工公式，需核对行变更引用后再更新')
                row = dict(zip(columns, raw)); row['日期'] = day(row['日期'])
                if row['日期'][:7] != sheet.title:
                    raise ValueError(f'商品总表跨月数据放错子表: {sheet.title}/{row["日期"]}')
                rows.append(row)
            months[sheet.title] = rows
    finally:
        wb.close()
    return months


def product_key(row):
    return tuple(str(row.get(k) or '') for k in ['日期', '店铺名称', '商品ID', 'SKU ID', '商品SKU'])


def product_sort(row):
    # Preserve the existing empty-burn-at-the-bottom convention within each month.
    return ('空烧推广费' in str(row.get('备注') or ''), *product_key(row))


def commit(output, payload, writer):
    deps = Path(os.environ.get('CODEX_WORKSPACE_DEPENDENCIES', str(Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies')))
    node = deps/('node/bin/node.exe' if os.name == 'nt' else 'node/bin/node')
    # Export to the same filesystem, replace only after a complete successful export.
    with tempfile.TemporaryDirectory(prefix='.日报更新-', dir=output.parent) as temporary:
        folder = Path(temporary)
        data = folder/'payload.json'; staged = folder/'report.xlsx'
        data.write_text(json.dumps(payload, ensure_ascii=False, default=str), encoding='utf-8')
        subprocess.run([str(node), str(SCRIPTS/writer), str(data), str(staged)], check=True)
        os.replace(staged, output)
        suffix = '.product.json' if writer == 'write_product_report.mjs' else '.store.json'
        os.replace(data, output.with_suffix(suffix))
        for extra in folder.glob('report.xlsx.*'):
            os.replace(extra, output.with_name(output.name+extra.name[len('report.xlsx'):]))


def product(output, columns, rows, dates, date_notes=None):
    output = Path(output).resolve(); dates = set(map(day, dates))
    incoming = []
    keys = set()
    for row in rows:
        row = dict(row); row['日期'] = day(row['日期'])
        if row['日期'] not in dates:
            raise ValueError('本次商品行超出请求日期，不更新历史记录')
        key = product_key(row)
        if key in keys:
            raise ValueError(f'本次商品明细键重复: {key}')
        keys.add(key); incoming.append(row)
    with locked(output):
        state_path=output.with_suffix('.product.json')
        state=json.loads(state_path.read_text(encoding='utf-8')) if state_path.exists() else {}
        previous_notes=state.get('date_notes', {})
        if any(dt in previous_notes and dt not in (date_notes or {}) for dt in dates):
            raise ValueError('本次日期已有确认的剔除/修订口径；请先按原规则处理数据，再提供date_notes或CLI --note，不能直接恢复原始订单')
        months = read_product(output, columns)
        affected = sorted({dt[:7] for dt in dates})
        sheets = []
        for month in affected:
            previous = months.get(month, [])
            merged = [r for r in previous if r['日期'] not in dates]
            merged.extend(r for r in incoming if r['日期'][:7] == month)
            merged.sort(key=product_sort)
            sheets.append({'name': month, 'rows': merged, 'oldRows': len(previous)})
            months[month] = merged
        payload = {'columns': columns, 'rows': [r for m in sorted(months) for r in months[m]],
                   'sheets': sheets, 'existingWorkbook': str(output) if output.exists() else None,
                   'replace_dates': sorted(dates), 'archive_mode': True,
                   'date_notes': {**previous_notes, **(date_notes or {})}}
        commit(output, payload, 'write_product_report.mjs')
    return output


def store_key(row):
    return row['日期'], row['平台'], row['店铺']


def cost_log(output, columns, rows, dates):
    path=Path(output).with_name(Path(output).stem+'_成本获取日志.csv')
    existing=[]
    if path.exists():
        with path.open(encoding='utf-8-sig',newline='') as f:
            existing=[r for r in csv.DictReader(f) if r.get('日期') not in dates]
    temporary=path.with_name(path.name+'.tmp')
    with temporary.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=columns);writer.writeheader();writer.writerows(existing+rows)
    os.replace(temporary,path)
    return path


def store(payload, directory, explicit_output=None):
    months = sorted({day(r['日期'])[:7] for r in payload['rows']})
    results = []
    for month in months:
        output = Path(explicit_output) if explicit_output and len(months) == 1 else Path(directory)/f'店铺日报_{month}.xlsx'
        output = output.resolve()
        with locked(output):
            previous = []
            state_path = output.with_suffix('.store.json')
            state = json.loads(state_path.read_text(encoding='utf-8')) if state_path.exists() else {}
            metadata = {store_key(r): r for r in state.get('rows', [])}
            if output.exists():
                from openpyxl import load_workbook
                wb = load_workbook(output, read_only=True, data_only=True)
                try:
                    sheet = wb['①平台日报']; sheet.reset_dimensions()
                    header = [v for v in next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))]
                    if header != payload['headers']:
                        raise ValueError('店铺月报表头不一致，不能自动覆盖')
                    for raw in sheet.iter_rows(min_row=4, values_only=True):
                        if not raw[0]:
                            continue
                        r = dict(zip(header, raw)); r['日期'] = day(r['日期'])
                        if r['日期'][:7] != month:
                            raise ValueError('店铺月报含其他月份数据，不能覆盖')
                        for k, v in metadata.get(store_key(r), {}).items():
                            if k.startswith('_'):
                                r[k] = v
                        # Preserve a known backend fallback for historical rows lacking visitors.
                        if r.get('访客数') is None and r.get('支付转化率') is not None:
                            r.setdefault('_后台支付转化率', r['支付转化率'])
                        previous.append(r)
                finally:
                    wb.close()
            incoming = [dict(r, 日期=day(r['日期'])) for r in payload['rows'] if day(r['日期'])[:7] == month]
            keys = [store_key(r) for r in incoming]
            if len(keys) != len(set(keys)):
                raise ValueError('本次店铺日汇总键重复')
            merged = {store_key(r): r for r in previous}
            for r in incoming:
                if merged.get(store_key(r),{}).get('_口径说明') and not r.get('_口径说明'):
                    raise ValueError('本次店铺日期已有确认修订；先恢复同日剔除/估算口径并携带_口径说明，再更新月报')
                for field in ['今日主要动作', '异常说明']:
                    if r.get(field) is None and store_key(r) in merged:
                        r[field] = merged[store_key(r)].get(field)
            merged.update({store_key(r): r for r in incoming})
            result = dict(payload, rows=[merged[k] for k in sorted(merged)], oldRows=len(previous),
                          existingWorkbook=str(output) if output.exists() else None, archive_mode=True,
                          report_note='本月累计日报；同一日期/平台/店铺重新生成时替换，不重复追加。历史估算和缺失口径按各日独立日志；今日动作/异常说明不自动填。')
            prefixes = tuple('|'.join(k)+'|' for k in keys)
            historical = {k:v for k,v in state.get('evidence', {}).items() if not k.startswith(prefixes)}
            result['evidence'] = {**historical, **payload.get('evidence', {})}
            logs = [*state.get('logs', []), *payload.get('logs', [])]
            result['logs'] = list({json.dumps(r,ensure_ascii=False,sort_keys=True):r for r in logs}.values())
            commit(output, result, 'write_store_report.mjs')
            log = output.with_name(output.stem+'_生成日志.json')
            log.write_text(json.dumps({'warnings': result['logs'], 'field_sources': result['evidence']}, ensure_ascii=False, indent=2), encoding='utf-8')
        results.append(output)
    return results


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('kind', choices=['product', 'store'])
    parser.add_argument('--input', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--date', action='append', default=[])
    parser.add_argument('--note',help='已核对并执行的同日修订口径，不是跳过原始计算的开关')
    args = parser.parse_args()
    data = json.loads(args.input.read_text(encoding='utf-8'))
    if args.kind == 'product':
        dates = args.date or sorted({day(r['日期']) for r in data['rows']})
        print(product(args.output, data['columns'], data['rows'], dates,
                      {dt:args.note for dt in dates} if args.note else None))
    else:
        print(store(data, args.output))
