"""Portable workbook I/O and V3 folder adapters; preserve original calculators."""
from pathlib import Path
from itertools import islice
from dataclasses import dataclass
import csv,json,os,subprocess

@dataclass
class WorkbookData:
    path: Path
    sheet_name: str
    rows: list
    total_rows: int=0
    truncated: bool=False

def read_workbook_sheets(path,max_rows=None,max_cols=None):
    path=Path(path)
    if path.suffix.lower()=='.csv':
        with path.open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
        return [WorkbookData(path,path.stem,rows[:max_rows] if max_rows else rows,len(rows),bool(max_rows and len(rows)>max_rows))]
    from openpyxl import load_workbook
    if path.suffix.lower()=='.xls':
        import xlrd
        wb=xlrd.open_workbook(str(path));out=[]
        for sheet in wb.sheets():
            rows=[sheet.row_values(i)[:max_cols] for i in range(min(sheet.nrows,max_rows or sheet.nrows))]
            out.append(WorkbookData(path,sheet.name,rows,sheet.nrows,bool(max_rows and sheet.nrows>max_rows)))
        return out
    wb=load_workbook(path,read_only=True,data_only=True);out=[]
    try:
        for sheet in wb.worksheets:
            sheet.reset_dimensions();iterator=sheet.iter_rows(max_col=max_cols,values_only=True)
            rows=[list(r) for r in (islice(iterator,max_rows+1) if max_rows else iterator)]
            truncated=bool(max_rows and len(rows)>max_rows)
            if truncated:rows=rows[:max_rows]
            # V3 exports contain repeated header names with values only in the first slot.
            if rows:
                headers=[str(v).strip() if v is not None else '' for v in rows[0]];positions={}
                for i,h in enumerate(headers):
                    if h:positions.setdefault(h,[]).append(i)
                for h,indices in positions.items():
                    if len(indices)<2:continue
                    for r in rows[1:]:
                        values=[r[i] for i in indices if i<len(r) and r[i] is not None and str(r[i]).strip()]
                        if len(set(map(str,values)))>1:raise ValueError(f'重复表头冲突 {path}/{sheet.title}/{h}')
                        if values:
                            for i in indices:
                                if i<len(r):r[i]=values[0]
            out.append(WorkbookData(path,sheet.title,rows,len(rows)+int(truncated),truncated))
    finally:wb.close()
    return out

def read_workbook(path,sheet_name=None,max_rows=None):
    sheets=read_workbook_sheets(path,max_rows)
    return next((s for s in sheets if s.sheet_name==sheet_name),sheets[0])

def install(module):
    module.read_workbook_sheets=read_workbook_sheets;module.read_workbook=read_workbook
    original_discover=module.discover_files
    def discover(root,subdir):
        current={'订单数据':'2.订单数据库','商品数据':'1.产品数据库','推广数据':'4.推广数据库','售后数据':'3.售后数据库'}.get(subdir)
        if (root/subdir).exists():return original_discover(root,subdir)
        if not current or not (root/current).exists():return []
        result=[]
        for p in (root/current).rglob('*'):
            if p.is_file() and p.suffix.lower() in {'.xlsx','.xls','.csv'} and not p.name.startswith('~$'):
                if subdir=='订单数据' and module.is_excluded_order_path(p):continue
                result.append(p)
        return sorted(result)
    module.discover_files=discover
    def write(path,columns,rows):
        path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
        data=path.with_suffix('.product.json');data.write_text(json.dumps({'columns':columns,'rows':rows},ensure_ascii=False,default=str),encoding='utf-8')
        deps=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies'
        node=deps/('node/bin/node.exe' if os.name=='nt' else 'node/bin/node')
        subprocess.run([str(node),str(Path(__file__).with_name('write_product_report.mjs')),str(data),str(path)],check=True)
    module.write_workbook=write
