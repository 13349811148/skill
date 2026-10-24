"""Store-day reporting, deliberately independent of product net-sales accounting."""
from __future__ import annotations
import argparse
import csv
import json
import math
import os
import subprocess
from datetime import date, datetime, timedelta
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP

SKILL = Path(__file__).resolve().parent.parent
HEADERS = ['日期','平台','店铺','填报人','访客数','曝光次数','点击次数','支付金额','支付订单数','支付买家数','退款金额','退款订单数','新客数','老客数','老客数占比','客单价','支付转化率','净成交额','推广花费','推广成交额','推广ROI','推广费比','店铺评分','投诉差评数','今日主要动作','异常说明']
DERIVED = {'客单价','支付转化率','净成交额','推广ROI','推广费比','新客数','老客数'}
ALWAYS_BLANK = {'今日主要动作','异常说明'}
ALIASES = {'店铺':['店铺','店铺名称','汇总_店铺名称'], '曝光次数':['推广曝光次数','推广曝光量'], '点击次数':['推广点击次数','推广点击量'], '推广成交额':['推广净成交额'], '老客数占比':['老客数占比','老买家占比']}

def estimate_customers(buyers, share):
    """User-approved estimate. Keep the source share rather than re-deriving it."""
    if buyers is None or share is None:return None,None
    count=Decimal(str(buyers));ratio=Decimal(str(share))
    if not count.is_finite() or count<0 or count!=count.to_integral_value():raise ValueError('支付买家数必须为非负整数')
    if not ratio.is_finite() or not Decimal(0)<=ratio<=Decimal(1):raise ValueError('老客数占比必须在0到1之间')
    old=int((count*ratio).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
    return int(count)-old,old

def number(value):
    if value is None or str(value).strip() in ('','-','--','n.a.','待补'): return None
    result = float(str(value).replace(',',''))
    if not math.isfinite(result) or result < 0: raise ValueError(f'无效或负的原始指标: {value}')
    return int(result) if result.is_integer() else result

def day(value):
    if isinstance(value,(date,datetime)): return value.strftime('%Y-%m-%d')
    return date.fromisoformat(str(value).strip()[:10]).isoformat()

def records(path):
    """Find headers in the first 20 rows, including hidden worksheets."""
    if path.suffix.lower()=='.csv':
        with path.open(encoding='utf-8-sig',newline='') as f:
            yield from csv.DictReader(f)
        return
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        for ws in wb.worksheets:
            ws.reset_dimensions(); iterator=ws.iter_rows(values_only=True); headers=None
            for n,row in enumerate(iterator,1):
                if n>20: break
                values=[str(v).strip() if v is not None else '' for v in row]
                if '日期' in values or '统计日期' in values:
                    headers=values;break
            if not headers: continue
            for row in iterator:
                if any(v is not None for v in row): yield dict(zip(headers,row))
    finally: wb.close()

def build_payload(root, dates, sources=(), config=None):
    config=config or json.loads((SKILL/'references/store_report_config.json').read_text(encoding='utf-8'))
    rows={}; logs=[]; evidence={}; selected={}
    def insert(record,source,fields,platform=None,shop=None):
        try: dt=day(record.get('日期') or record.get('统计日期'))
        except (ValueError,TypeError): return
        if dt not in dates:return
        platform=record.get('平台') or platform
        shop=next((record.get(k) for k in ALIASES['店铺'] if record.get(k)),shop)
        shop=config.get('shop_aliases',{}).get(str(shop),shop)
        if not platform or not shop:
            logs.append({'source':str(source),'date':dt,'message':'缺少明确平台或店铺，未导入'});return
        key=(dt,str(platform),str(shop))
        row=rows.setdefault(key,{h:None for h in HEADERS});row.update(zip(HEADERS[:3],key))
        row['填报人']=config.get('reporters',{}).get(str(shop))
        rank=source.stat().st_mtime_ns
        for field,aliases in fields.items():
            raw=next((record.get(a) for a in aliases if record.get(a) is not None and str(record.get(a)).strip()),None)
            value=number(str(Decimal(str(raw).strip().rstrip('%'))/100)) if field=='老客数占比' and raw is not None and str(raw).strip().endswith('%') else number(raw)
            if value is None:continue
            if field=='老客数占比' and value>1:raise ValueError(f'{source}: 老客数占比必须在0到1之间')
            if field in {'访客数','曝光次数','点击次数','支付订单数','支付买家数','退款订单数','投诉差评数'} and int(value)!=value:
                raise ValueError(f'{source}: {field}不是整数')
            slot=(*key,field)
            if slot in selected and rank==selected[slot] and row[field]!=value:
                raise ValueError(f'同一来源包含冲突的店铺日指标: {slot}；不能直接累加人数或整店汇总')
            if slot not in selected or rank>selected[slot]:
                row[field]=value;selected[slot]=rank;evidence['|'.join(slot)]=str(source)
    # Explicit daily shop metrics: accept no ambiguous product-order/net-sales substitutes.
    found=list(sources)
    for folder in ['店铺日报输入','店铺数据','7.店铺数据库']:
        if (root/folder).exists():found.extend(p for p in (root/folder).rglob('*') if p.suffix.lower() in {'.xlsx','.csv'} and not p.name.startswith('~$'))
    fields={h:ALIASES.get(h,[h]) for h in HEADERS[4:] if h not in DERIVED|ALWAYS_BLANK}
    for path in sorted(set(map(Path,found))):
        for record in records(path):insert(record,path,fields)
    # V3 backend archive is an explicit daily source, not an order-net-sales substitute.
    backend=root/'7.后台数据库'
    if backend.exists():
        for path in sorted(backend.rglob('*.xlsx')):
            if path.name.startswith('~$'):continue
            meta=path.with_suffix('.archive.json')
            if not meta.exists():continue
            metadata=json.loads(meta.read_text(encoding='utf-8-sig'))
            backend_shop=metadata.get('shop')
            canonical=config.get('shop_aliases',{}).get(str(backend_shop))
            if not canonical:
                logs.append({'source':str(path),'message':'后台店名未确认映射，未与订单/推广合并'});continue
            for record in records(path):
                insert(record,path,{'访客数':['店铺访客数（人）','店铺访客数','访客数'],'支付金额':['成交金额（元）'],'支付订单数':['订单数'],'支付买家数':['成交买家数'],'老客数占比':['老买家占比'],'退款金额':['退款金额（元）'],'退款订单数':['退款单数'],'店铺评分':['综合体验星级'],'投诉差评数':['新增差评（条）']},'拼多多',canonical)
                try:dt=day(record.get('日期'))
                except (ValueError,TypeError):continue
                if dt not in dates:continue
                key=(dt,'拼多多',canonical)
                raw_rate=record.get('转化率')
                if raw_rate is not None:
                    rate=number(raw_rate)
                    if rate is not None and rate>1:raise ValueError(f'后台转化率需为0到1的数值 {path}/{dt}')
                    old=rows[key].get('_后台转化率来源时间',-1)
                    if path.stat().st_mtime_ns>=old:
                        rows[key]['_后台支付转化率']=rate;rows[key]['_后台转化率来源时间']=path.stat().st_mtime_ns
                        evidence['|'.join((*key,'支付转化率'))]=str(path)+'（后台原始转化率，访客缺失时使用；不反推访客）'
            logs.append({'source':str(path),'message':'后台日成交金额/订单数/买家数对应未扣退款支付口径；退款采用后台当日退款指标；评分=综合体验星级，投诉差评列仅填新增差评（1至3星，不含默认评价）；老客占比取后台原值，新老买家人数按用户确认规则重新推算。'})
    # Account daily exports contain total promotion metrics, without a product ID.
    for folder in ['4.推广数据库','推广数据']:
        if not (root/folder).exists():continue
        for path in sorted((root/folder).rglob('*.xlsx')):
            if path.name.startswith('~$') or '账户每日汇总' not in str(path):continue
            if any(x in str(path).lower() for x in ['备份','backup','副本']):
                raise ValueError(f'推广输入包含备份，需移出输入目录: {path}')
            meta=path.with_suffix(path.suffix+'.archive.json')
            metadata=json.loads(meta.read_text(encoding='utf-8-sig')) if meta.exists() else {}
            shop=metadata.get('shop')
            platform=metadata.get('platform') or ('拼多多' if shop and '拼多多' in shop else None)
            for record in records(path):
                if any(record.get(k) for k in ['商品ID','商品id','归档商品ID']):
                    raise ValueError(f'账户每日汇总中发现商品明细，不能直接当店铺汇总: {path}')
                insert(record,path,{'曝光次数':['曝光量'],'点击次数':['点击量'],'推广花费':['总花费(元)','成交花费(元)'],'推广成交额':['净交易额(元)']},platform,shop)
    configured=config.get('shops',[])
    for shop in configured:
        for dt in dates:
            key=(dt,shop['平台'],shop['店铺'])
            row=rows.setdefault(key,{h:None for h in HEADERS});row.update(zip(HEADERS[:3],key));row['填报人']=config.get('reporters',{}).get(shop['店铺'])
    for row in rows.values():
        for field in ALWAYS_BLANK:row[field]=None
        row['新客数'],row['老客数']=estimate_customers(row.get('支付买家数'),row.get('老客数占比'))
        if row['老客数'] is not None:
            key=(row['日期'],row['平台'],row['店铺'])
            for field in ['新客数','老客数']:
                evidence['|'.join((*key,field))]='推算值（用户确认）：老客数=ROUND(支付买家数×后台老买家占比,0)，新客数=支付买家数−老客数；原始比例来源='+evidence.get('|'.join((*key,'老客数占比')),'未知')
            logs.append({'date':row['日期'],'shop':row['店铺'],'message':f'新老客人数为推算值（用户已确认）；支付买家数={row["支付买家数"]}，后台老客占比={row["老客数占比"]}，新客数={row["新客数"]}，老客数={row["老客数"]}；占比保留原值，不用取整后人数反算。'})
        visitors=row.get('访客数');buyers=row.get('支付买家数');raw_rate=row.get('_后台支付转化率')
        if visitors is not None and buyers is not None and visitors>0 and raw_rate is not None:
            calculated=buyers/visitors
            if abs(calculated-raw_rate)>0.00005001:
                logs.append({'date':row['日期'],'shop':row['店铺'],'message':f'支付转化率与后台显示值超出0.005个百分点：按原始买家/访客计算={calculated:.8%}，后台原值={raw_rate:.8%}；保留公式结果，请核对。'})
            else:
                logs.append({'date':row['日期'],'shop':row['店铺'],'message':'支付买家数÷原始店铺访客数与后台转化率按两位百分数核对一致。'})
        # Missing is not zero. Source data may legitimately provide a numeric zero.
        missing=[h for h in HEADERS[4:24] if h not in DERIVED|ALWAYS_BLANK and row[h] is None]
        if missing:logs.append({'date':row['日期'],'platform':row['平台'],'shop':row['店铺'],'message':'缺少数据，保持空白: '+','.join(missing)})
        if row['支付金额'] is not None and row['退款金额'] is not None and row['支付金额']-row['退款金额']<0:
            logs.append({'date':row['日期'],'shop':row['店铺'],'message':'全店净成交额为负，保留净成交额；推广ROI和推广费比留空'})
    return {'headers':HEADERS,'rows':[rows[k] for k in sorted(rows)],'logs':logs,'evidence':evidence}

def main():
    p=argparse.ArgumentParser()
    p.add_argument('--database-root',default=str(Path.home()/'Desktop/运营数据库'))
    p.add_argument('--input',action='append',default=[])
    p.add_argument('--config',type=Path)
    p.add_argument('--date',action='append');p.add_argument('--start-date');p.add_argument('--end-date')
    p.add_argument('--output',type=Path);p.add_argument('--payload-only',action='store_true')
    a=p.parse_args();dates=set()
    for item in a.date or []:dates.update(day(v) for v in item.split(','))
    if bool(a.start_date)!=bool(a.end_date):p.error('开始/结束日期须同时提供')
    if a.start_date:
        start=date.fromisoformat(a.start_date);end=date.fromisoformat(a.end_date)
        if end<start:p.error('结束日期早于开始日期')
        dates.update((start+timedelta(days=i)).isoformat() for i in range((end-start).days+1))
    if not dates:p.error('必须指定统计日期')
    root=Path(a.database_root);output=a.output or root/'数据报表输出/店铺维度日报.xlsx'
    payload=build_payload(root,dates,a.input,json.loads(a.config.read_text(encoding='utf-8')) if a.config else None)
    output.parent.mkdir(parents=True,exist_ok=True)
    payload_path=output.with_suffix('.store.json');payload_path.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    log=output.with_name(output.stem+'_生成日志.json');log.write_text(json.dumps({'warnings':payload['logs'],'field_sources':payload['evidence']},ensure_ascii=False,indent=2),encoding='utf-8')
    if not a.payload_only:
        deps=Path(os.environ.get('CODEX_WORKSPACE_DEPENDENCIES',str(Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies')))
        node=deps/'node/bin/node.exe' if os.name=='nt' else deps/'node/bin/node'
        # Script lives next to its bundled node_modules junction, created at install.
        subprocess.run([str(node),str(Path(__file__).with_name('write_store_report.mjs')),str(payload_path),str(output)],check=True)
    print(f'店铺日报: {output}; 行数={len(payload["rows"])}; 缺失/异常记录={len(payload["logs"])}; 日志={log}')
    return 0

if __name__=='__main__':raise SystemExit(main())
