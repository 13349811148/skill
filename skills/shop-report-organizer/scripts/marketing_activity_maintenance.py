#!/usr/bin/env python
"""Plan and apply safe maintenance for the operations marketing-activity workbook."""

from __future__ import annotations

import datetime as dt
import json
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable


MARKETING_FOLDER_NAME = "营销活动监控"
BACKUP_FOLDER_NAME = "备份"
SUPPORTED_EXTENSIONS = {".xlsx", ".xls"}

COMPACT_HEADERS = [
    "样式ID",
    "商品ID",
    "品牌",
    "所属运营",
    "平台",
    "店铺名称",
    "大类",
    "商品SKU",
    "产品成本",
    "活动价",
    "报名价",
    "促销机制",
]

TRACKING_HEADERS = [
    "品牌",
    "所属运营",
    "平台",
    "店铺名称",
    "大类",
    "商品ID",
    "样式ID",
    "商品SKU",
    "产品成本",
    "营销活动ID",
    "活动名称",
    "活动价",
    "报名价",
    "其他店铺优惠",
    "单件预估实收",
    "多件折扣",
    "多件预估实收",
    "排期开始日期",
    "排期结束日期",
    "补贴金额",
    "单件预估实收（补贴后）",
    "单件预估毛利",
    "单件预估毛利率",
    "多件预估实收（补贴后）",
    "多件预估毛利",
    "多件预估毛利率",
]

BASE_FIELDS = ["店铺名称", "商品ID", "样式ID", "商品SKU"]
HISTORICAL_FIELDS = ["品牌", "所属运营", "平台", "大类"]
AUTOMATIC_FIELDS = [*BASE_FIELDS, *HISTORICAL_FIELDS, "产品成本"]
COMPACT_MANUAL_FIELDS = ["活动价", "报名价", "促销机制"]
TRACKING_MANUAL_FIELDS = [
    "营销活动ID",
    "活动名称",
    "活动价",
    "报名价",
    "单件预估实收",
    "多件预估实收",
]
TRACKING_SILENT_FIELDS = ["其他店铺优惠", "多件折扣", "排期开始日期", "排期结束日期"]
TRACKING_FORMULA_FIELDS = [
    "补贴金额",
    "单件预估实收（补贴后）",
    "单件预估毛利",
    "单件预估毛利率",
    "多件预估实收（补贴后）",
    "多件预估毛利",
    "多件预估毛利率",
]
TRACKING_FORMULAS_R1C1 = {
    "补贴金额": '=IFERROR(RC[-7]-RC[-8],"")',
    "单件预估实收（补贴后）": '=IFERROR(IF(RC[-6]+RC[-1]=0,"",RC[-6]+RC[-1]),"")',
    "单件预估毛利": '=IFERROR(RC[-1]-RC[-13],"")',
    "单件预估毛利率": '=IFERROR(RC[-1]/RC[-2],"")',
    "多件预估实收（补贴后）": '=IFERROR(IF(RC[-7]+RC[-4]=0,"",RC[-4]+RC[-7]),"")',
    "多件预估毛利": '=IFERROR(RC[-1]-RC[-16],"")',
    "多件预估毛利率": '=IFERROR(RC[-1]/RC[-2],"")',
}


@dataclass
class WorkbookSchema:
    path: Path
    layout: str
    sheet_name: str
    data_start_row: int
    columns: dict[str, int]
    max_column: int

    def as_dict(self) -> dict[str, object]:
        return {
            "path": str(self.path),
            "layout": self.layout,
            "sheet_name": self.sheet_name,
            "data_start_row": self.data_start_row,
            "columns": self.columns,
            "max_column": self.max_column,
        }


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def normalize_header(value: object) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", normalize_text(value))).casefold()


def normalize_key(value: object) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", normalize_text(value))).casefold()


def normalize_product_id(value: object) -> str:
    text = normalize_text(value)
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    digits = re.sub(r"\D", "", text)
    return digits if len(digits) >= 6 else ""


def json_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (dt.date, dt.datetime, dt.time)):
        return value.isoformat(sep=" ")
    return normalize_text(value)


def value_token(value: object) -> str:
    if value is None or normalize_text(value) == "":
        return ""
    if isinstance(value, bool):
        return f"bool:{value}"
    if isinstance(value, (int, float)):
        try:
            number = float(value)
        except (TypeError, ValueError):
            pass
        else:
            return f"num:{number:.12g}"
    text = normalize_text(value)
    if re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text):
        try:
            return f"num:{float(text):.12g}"
        except ValueError:
            pass
    return f"text:{normalize_key(text)}"


def values_equal(left: object, right: object) -> bool:
    return value_token(left) == value_token(right)


def canonical_header_map(values: Iterable[object]) -> dict[str, int]:
    known = {normalize_header(header): header for header in {*COMPACT_HEADERS, *TRACKING_HEADERS}}
    columns: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        normalized_value = normalize_header(value)
        field = known.get(normalized_value)
        if field is None and normalized_value:
            annotated_matches = [
                (normalized_header, candidate)
                for normalized_header, candidate in known.items()
                if normalized_value.startswith(normalized_header)
            ]
            if annotated_matches:
                field = max(annotated_matches, key=lambda item: len(item[0]))[1]
        if field and field not in columns:
            columns[field] = index
    return columns


def classify_header_rows(
    path: Path,
    sheet_name: str,
    rows: list[list[object]],
    physical_rows: list[int],
) -> WorkbookSchema | None:
    compact_required = {"商品ID", "样式ID", "商品SKU", "产品成本", "活动价", "报名价", "促销机制"}
    tracking_required = {
        "商品ID",
        "样式ID",
        "商品SKU",
        "营销活动ID",
        "单件预估实收",
        "多件预估实收",
        "补贴金额",
    }
    for offset, row in enumerate(rows):
        compact_columns = canonical_header_map(row)
        if compact_required.issubset(compact_columns) and "营销活动ID" not in compact_columns:
            return WorkbookSchema(
                path=path,
                layout="compact",
                sheet_name=sheet_name,
                data_start_row=physical_rows[offset] + 1,
                columns=compact_columns,
                max_column=max(compact_columns.values()),
            )
        if offset + 1 >= len(rows):
            continue
        upper = row
        lower = rows[offset + 1]
        width = max(len(upper), len(lower))
        combined = [
            lower[index] if index < len(lower) and normalize_text(lower[index]) else (
                upper[index] if index < len(upper) else ""
            )
            for index in range(width)
        ]
        tracking_columns = canonical_header_map(combined)
        if tracking_required.issubset(tracking_columns):
            # The confirmed tracking layout has a fixed A:Z physical schema.
            # Some formula subheaders abbreviate repeated "单件/多件" labels,
            # so use the validated physical order after signature detection.
            tracking_columns = {
                field: index
                for index, field in enumerate(TRACKING_HEADERS, start=1)
            }
            return WorkbookSchema(
                path=path,
                layout="tracking",
                sheet_name=sheet_name,
                data_start_row=physical_rows[offset + 1] + 1,
                columns=tracking_columns,
                max_column=len(TRACKING_HEADERS),
            )
    return None


def classify_xlsx(path: Path) -> WorkbookSchema | None:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment specific
        raise RuntimeError("openpyxl is required for marketing activity maintenance") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        for worksheet in workbook.worksheets:
            rows = [
                [cell.value for cell in row]
                for row in worksheet.iter_rows(min_row=1, max_row=6, max_col=40)
            ]
            schema = classify_header_rows(
                path,
                worksheet.title,
                rows,
                list(range(1, len(rows) + 1)),
            )
            if schema:
                return schema
    finally:
        workbook.close()
    return None


def split_xls_marker_rows(
    rows: list[list[str]],
    marker: str,
) -> list[tuple[str, list[list[str]], list[int]]]:
    sheets: list[tuple[str, list[list[str]], list[int]]] = []
    sheet_name = ""
    sheet_rows: list[list[str]] = []
    physical_rows: list[int] = []
    next_physical_row = 1
    for row in rows:
        if row and row[0] == marker:
            if sheet_name:
                sheets.append((sheet_name, sheet_rows, physical_rows))
            sheet_name = normalize_text(row[1]) if len(row) > 1 else "Sheet1"
            sheet_rows = []
            physical_rows = []
            next_physical_row = 1
            continue
        if not sheet_name:
            sheet_name = "Sheet1"
        sheet_rows.append(row)
        physical_rows.append(next_physical_row)
        next_physical_row += 1
    if sheet_name:
        sheets.append((sheet_name, sheet_rows, physical_rows))
    return sheets


def classify_xls(
    path: Path,
    xls_reader: Callable[..., list[list[str]]],
    marker: str,
) -> WorkbookSchema | None:
    rows = xls_reader(
        path,
        12,
        40,
        all_worksheets=True,
        include_sheet_markers=True,
    )
    for sheet_name, sheet_rows, physical_rows in split_xls_marker_rows(rows, marker):
        schema = classify_header_rows(path, sheet_name, sheet_rows[:6], physical_rows[:6])
        if schema:
            return schema
    return None


def discover_schema(
    database_root: Path,
    xls_reader: Callable[..., list[list[str]]] | None,
    xls_sheet_marker: str,
) -> tuple[WorkbookSchema | None, list[dict[str, str]], list[WorkbookSchema]]:
    folder = database_root / MARKETING_FOLDER_NAME
    issues: list[dict[str, str]] = []
    if not folder.exists():
        issues.append(
            {
                "kind": "missing_marketing_folder",
                "message": f"未找到运营数据库中的“{MARKETING_FOLDER_NAME}”文件夹，营销活动表未维护。",
            }
        )
        return None, issues, []
    schemas: list[WorkbookSchema] = []
    for path in sorted(folder.iterdir()):
        if (
            not path.is_file()
            or path.name.startswith("~$")
            or path.suffix.lower() not in SUPPORTED_EXTENSIONS
        ):
            continue
        try:
            if path.suffix.lower() == ".xlsx":
                schema = classify_xlsx(path)
            elif xls_reader is not None:
                schema = classify_xls(path, xls_reader, xls_sheet_marker)
            else:
                schema = None
        except Exception as exc:
            issues.append(
                {
                    "kind": "workbook_inspection_error",
                    "message": f"检查营销活动候选文件失败：{path.name}（{exc}）。",
                }
            )
            continue
        if schema:
            schemas.append(schema)
    layouts = {schema.layout for schema in schemas}
    if len(layouts) > 1:
        names = "、".join(schema.path.name for schema in schemas)
        issues.append(
            {
                "kind": "multiple_layouts",
                "message": f"运营数据库中同时存在简版和跟踪版营销活动表（{names}），本次未自动选择或维护。",
            }
        )
        return None, issues, schemas
    if not schemas:
        issues.append(
            {
                "kind": "missing_layout",
                "message": "运营数据库中未找到可识别的简版或跟踪版营销活动表，本次未维护。",
            }
        )
        return None, issues, []
    selected = max(
        schemas,
        key=lambda item: (
            item.path.stat().st_mtime_ns,
            item.path.suffix.lower() == ".xlsx",
            item.path.name.casefold(),
        ),
    )
    return selected, issues, schemas


def row_values_from_cells(
    cells: Iterable[object],
    schema: WorkbookSchema,
) -> dict[str, object]:
    values = list(cells)
    result: dict[str, object] = {}
    for field, column in schema.columns.items():
        value = values[column - 1] if column - 1 < len(values) else ""
        if hasattr(value, "value"):
            value = value.value
        result[field] = json_value(value)
    return result


def scan_xlsx_rows(schema: WorkbookSchema) -> list[dict[str, object]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment specific
        raise RuntimeError("openpyxl is required for marketing activity maintenance") from exc
    workbook = openpyxl.load_workbook(schema.path, read_only=True, data_only=False)
    try:
        worksheet = workbook[schema.sheet_name]
        records: list[dict[str, object]] = []
        blank_base_streak = 0
        for row_number, cells in enumerate(
            worksheet.iter_rows(
                min_row=schema.data_start_row,
                max_col=schema.max_column,
            ),
            start=schema.data_start_row,
        ):
            values = row_values_from_cells(cells, schema)
            if not any(normalize_text(values.get(field, "")) for field in BASE_FIELDS):
                blank_base_streak += 1
                if blank_base_streak >= 50:
                    break
                continue
            blank_base_streak = 0
            records.append({"row": row_number, "values": values})
        return records
    finally:
        workbook.close()


def scan_xls_rows(
    schema: WorkbookSchema,
    xls_reader: Callable[..., list[list[str]]],
    marker: str,
) -> list[dict[str, object]]:
    rows = xls_reader(
        schema.path,
        65_536,
        max(40, schema.max_column),
        all_worksheets=True,
        include_sheet_markers=True,
    )
    records: list[dict[str, object]] = []
    for sheet_name, sheet_rows, physical_rows in split_xls_marker_rows(rows, marker):
        if sheet_name != schema.sheet_name:
            continue
        for row, physical_row in zip(sheet_rows, physical_rows):
            if physical_row < schema.data_start_row:
                continue
            values = row_values_from_cells(row, schema)
            if any(normalize_text(values.get(field, "")) for field in BASE_FIELDS):
                records.append({"row": physical_row, "values": values})
        break
    return records


def scan_existing_rows(
    schema: WorkbookSchema,
    xls_reader: Callable[..., list[list[str]]] | None,
    xls_sheet_marker: str,
) -> list[dict[str, object]]:
    if schema.path.suffix.lower() == ".xlsx":
        return scan_xlsx_rows(schema)
    if xls_reader is None:
        raise RuntimeError("legacy .xls marketing workbook requires the Excel reader")
    return scan_xls_rows(schema, xls_reader, xls_sheet_marker)


def detect_cost_header(row: list[object]) -> tuple[int, int] | None:
    code_column = next(
        (index for index, value in enumerate(row) if normalize_header(value) == normalize_header("商品编码")),
        None,
    )
    cost_column = next(
        (
            index
            for index, value in enumerate(row)
            if "成本价" in normalize_header(value)
        ),
        None,
    )
    if code_column is None or cost_column is None:
        return None
    return code_column, cost_column


def iter_cost_rows_xlsx(path: Path) -> Iterable[tuple[object, object]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment specific
        raise RuntimeError("openpyxl is required for marketing activity maintenance") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            header: tuple[int, int] | None = None
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=1, max_col=40, values_only=True),
                start=1,
            ):
                values = list(row)
                if row_number <= 12 and header is None:
                    header = detect_cost_header(values)
                    if header is not None:
                        continue
                if header is None:
                    if row_number >= 12:
                        break
                    continue
                code_column, cost_column = header
                code = values[code_column] if code_column < len(values) else ""
                cost = values[cost_column] if cost_column < len(values) else ""
                if normalize_text(code):
                    yield code, cost
    finally:
        workbook.close()


def iter_cost_rows_xls(
    path: Path,
    xls_reader: Callable[..., list[list[str]]],
    marker: str,
) -> Iterable[tuple[object, object]]:
    rows = xls_reader(
        path,
        65_536,
        40,
        all_worksheets=True,
        include_sheet_markers=True,
    )
    for _sheet_name, sheet_rows, _physical_rows in split_xls_marker_rows(rows, marker):
        header: tuple[int, int] | None = None
        for row_number, row in enumerate(sheet_rows, start=1):
            if row_number <= 12 and header is None:
                header = detect_cost_header(row)
                if header is not None:
                    continue
            if header is None:
                if row_number >= 12:
                    break
                continue
            code_column, cost_column = header
            code = row[code_column] if code_column < len(row) else ""
            cost = row[cost_column] if cost_column < len(row) else ""
            if normalize_text(code):
                yield code, cost


def load_cost_map(
    folder: Path,
    selected_workbook: Path,
    xls_reader: Callable[..., list[list[str]]] | None,
    xls_sheet_marker: str,
) -> tuple[dict[str, object], set[str]]:
    raw_values: dict[str, dict[str, object]] = defaultdict(dict)
    for path in sorted(folder.iterdir()):
        if (
            not path.is_file()
            or path == selected_workbook
            or path.name.startswith("~$")
            or path.suffix.lower() not in SUPPORTED_EXTENSIONS
        ):
            continue
        try:
            if path.suffix.lower() == ".xlsx":
                cost_rows = iter_cost_rows_xlsx(path)
            elif xls_reader is not None:
                cost_rows = iter_cost_rows_xls(path, xls_reader, xls_sheet_marker)
            else:
                continue
            for sku, cost in cost_rows:
                sku_key = normalize_key(sku)
                token = value_token(cost)
                if sku_key and token:
                    raw_values[sku_key].setdefault(token, json_value(cost))
        except Exception:
            continue
    unique: dict[str, object] = {}
    conflicts: set[str] = set()
    for sku_key, values in raw_values.items():
        if len(values) == 1:
            unique[sku_key] = next(iter(values.values()))
        elif values:
            conflicts.add(sku_key)
    return unique, conflicts


def load_history(
    rows: list[dict[str, object]],
) -> tuple[dict[str, dict[str, object]], dict[str, set[str]]]:
    raw: dict[str, dict[str, dict[str, object]]] = defaultdict(
        lambda: defaultdict(dict)
    )
    for record in rows:
        values = record["values"]
        sku_key = normalize_key(values.get("商品SKU", ""))
        if not sku_key:
            continue
        for field in HISTORICAL_FIELDS:
            value = values.get(field, "")
            token = value_token(value)
            if token:
                raw[sku_key][field].setdefault(token, json_value(value))
    unique: dict[str, dict[str, object]] = defaultdict(dict)
    conflicts: dict[str, set[str]] = defaultdict(set)
    for sku_key, fields in raw.items():
        for field, values in fields.items():
            if len(values) == 1:
                unique[sku_key][field] = next(iter(values.values()))
            elif values:
                conflicts[sku_key].add(field)
    return dict(unique), dict(conflicts)


def normalize_candidates(
    candidates: list[dict[str, object]],
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    grouped: dict[tuple[str, str, str], dict[str, set[str]]] = {}
    issues: list[dict[str, str]] = []
    for candidate in candidates:
        shop_name = normalize_text(candidate.get("shop_name", ""))
        product_id = normalize_product_id(candidate.get("product_id", ""))
        style_id = normalize_text(candidate.get("style_id", ""))
        sku = normalize_text(candidate.get("sku", ""))
        source = normalize_text(candidate.get("source_name", ""))
        if not product_id:
            continue
        key = (shop_name, product_id, style_id)
        item = grouped.setdefault(key, {"skus": set(), "sources": set(), "titles": set()})
        if sku:
            item["skus"].add(sku)
        if source:
            item["sources"].add(source)
        title = normalize_text(candidate.get("title", ""))
        if title:
            item["titles"].add(title)
    normalized: list[dict[str, str]] = []
    for (shop_name, product_id, style_id), item in sorted(grouped.items()):
        if not shop_name:
            issues.append(
                {
                    "kind": "missing_shop",
                    "product_id": product_id,
                    "style_id": style_id,
                    "sku": "",
                    "message": f"商品ID {product_id} 未确定店铺，营销活动表未新增该记录。",
                }
            )
            continue
        if not style_id:
            issues.append(
                {
                    "kind": "missing_style_id",
                    "product_id": product_id,
                    "style_id": "",
                    "sku": "",
                    "message": f"商品ID {product_id} 未找到样式ID，营销活动表未新增该记录，请人工维护。",
                }
            )
            continue
        if len(item["skus"]) > 1:
            sku_text = "、".join(sorted(item["skus"]))
            issues.append(
                {
                    "kind": "conflicting_sku",
                    "product_id": product_id,
                    "style_id": style_id,
                    "sku": sku_text,
                    "message": (
                        f"商品ID {product_id}、样式ID {style_id} 对应多个商家编码-规格维度"
                        f"（{sku_text}），营销活动表未自动写入，请人工核对。"
                    ),
                }
            )
            continue
        normalized.append(
            {
                "shop_name": shop_name,
                "product_id": product_id,
                "style_id": style_id,
                "sku": next(iter(item["skus"]), ""),
                "source_name": "；".join(sorted(item["sources"])),
                "title": "；".join(sorted(item["titles"])),
            }
        )
    return normalized, issues


def issue_for(
    kind: str,
    candidate: dict[str, str],
    message: str,
) -> dict[str, str]:
    return {
        "kind": kind,
        "product_id": candidate["product_id"],
        "style_id": candidate["style_id"],
        "sku": candidate["sku"],
        "message": message,
    }


def find_target_row(
    candidate: dict[str, str],
    rows: list[dict[str, object]],
) -> tuple[dict[str, object] | None, str]:
    product_id = candidate["product_id"]
    style_key = normalize_key(candidate["style_id"])
    sku_key = normalize_key(candidate["sku"])
    product_style_rows = [
        row
        for row in rows
        if normalize_product_id(row["values"].get("商品ID", "")) == product_id
        and normalize_key(row["values"].get("样式ID", "")) == style_key
    ]
    if sku_key:
        exact = [
            row
            for row in product_style_rows
            if normalize_key(row["values"].get("商品SKU", "")) == sku_key
        ]
        if len(exact) == 1:
            return exact[0], ""
        if len(exact) > 1:
            return None, "duplicate_exact_rows"
        blank_sku = [
            row for row in product_style_rows if not normalize_key(row["values"].get("商品SKU", ""))
        ]
        nonblank_other = [
            row for row in product_style_rows if normalize_key(row["values"].get("商品SKU", ""))
        ]
        if len(blank_sku) == 1 and not nonblank_other:
            return blank_sku[0], ""
        if product_style_rows:
            return None, "existing_sku_conflict"
        return None, ""
    if len(product_style_rows) == 1:
        return product_style_rows[0], ""
    if len(product_style_rows) > 1:
        return None, "ambiguous_without_sku"
    return None, ""


def desired_values(
    candidate: dict[str, str],
    history: dict[str, dict[str, object]],
    history_conflicts: dict[str, set[str]],
    costs: dict[str, object],
    cost_conflicts: set[str],
    current_values: dict[str, object],
) -> tuple[dict[str, object], list[dict[str, str]]]:
    desired: dict[str, object] = {
        "店铺名称": candidate["shop_name"],
        "商品ID": candidate["product_id"],
        "样式ID": candidate["style_id"],
    }
    issues: list[dict[str, str]] = []
    sku = candidate["sku"]
    if sku:
        desired["商品SKU"] = sku
        sku_key = normalize_key(sku)
        historical_missing: list[str] = []
        historical_conflicting: list[str] = []
        for field in HISTORICAL_FIELDS:
            if field in history_conflicts.get(sku_key, set()):
                if not normalize_text(current_values.get(field, "")):
                    historical_conflicting.append(field)
            elif field in history.get(sku_key, {}):
                desired[field] = history[sku_key][field]
            elif not normalize_text(current_values.get(field, "")):
                historical_missing.append(field)
        if historical_conflicting:
            issues.append(
                issue_for(
                    "historical_conflict",
                    candidate,
                    (
                        f"商品ID {candidate['product_id']}、商品SKU {sku} 的历史"
                        f"{'、'.join(historical_conflicting)}存在冲突，未自动填写，请人工维护。"
                    ),
                )
            )
        if historical_missing:
            issues.append(
                issue_for(
                    "historical_missing",
                    candidate,
                    (
                        f"商品ID {candidate['product_id']}、商品SKU {sku} 没有可唯一匹配的历史"
                        f"{'、'.join(historical_missing)}，请人工维护。"
                    ),
                )
            )
        if sku_key in cost_conflicts and not normalize_text(current_values.get("产品成本", "")):
            issues.append(
                issue_for(
                    "cost_conflict",
                    candidate,
                    f"商品ID {candidate['product_id']}、商品SKU {sku} 在成本表中对应多个成本，未自动填写产品成本。",
                )
            )
        elif sku_key in costs:
            desired["产品成本"] = costs[sku_key]
        elif not normalize_text(current_values.get("产品成本", "")):
            issues.append(
                issue_for(
                    "cost_missing",
                    candidate,
                    f"商品ID {candidate['product_id']}、商品SKU {sku} 在成本表中未找到产品成本，请人工维护。",
                )
            )
    elif not normalize_text(current_values.get("商品SKU", "")):
        issues.append(
            issue_for(
                "missing_sku",
                candidate,
                (
                    f"商品ID {candidate['product_id']}、样式ID {candidate['style_id']} "
                    "未找到商家编码-规格维度（商品SKU）；已保留可确定字段，请人工维护商品SKU及其成本/历史字段。"
                ),
            )
        )
    return desired, issues


def prepare_marketing_activity_plan(
    database_root: Path,
    candidates: list[dict[str, object]],
    *,
    xls_reader: Callable[..., list[list[str]]] | None = None,
    xls_sheet_marker: str = "__CODEX_ORDER_SHEET__",
) -> dict[str, object]:
    normalized_candidates, candidate_issues = normalize_candidates(candidates)
    plan: dict[str, object] = {
        "status": "no_candidates" if not normalized_candidates and not candidate_issues else "pending",
        "database_root": str(database_root),
        "schema": None,
        "candidates": normalized_candidates,
        "writes": [],
        "issues": candidate_issues,
        "manual_items": [],
        "available_schemas": [],
    }
    if not normalized_candidates:
        return plan

    schema, discovery_issues, all_schemas = discover_schema(
        database_root,
        xls_reader,
        xls_sheet_marker,
    )
    plan["issues"].extend(discovery_issues)
    plan["available_schemas"] = [item.as_dict() for item in all_schemas]
    if schema is None:
        plan["status"] = "skipped"
        return plan
    plan["schema"] = schema.as_dict()

    rows = scan_existing_rows(schema, xls_reader, xls_sheet_marker)
    history, history_conflicts = load_history(rows)
    costs, cost_conflicts = load_cost_map(
        schema.path.parent,
        schema.path,
        xls_reader,
        xls_sheet_marker,
    )
    max_existing_row = max(
        (int(record["row"]) for record in rows),
        default=schema.data_start_row - 1,
    )
    next_row = max(max_existing_row + 1, schema.data_start_row)
    template_row = max_existing_row if max_existing_row >= schema.data_start_row else schema.data_start_row
    manual_fields = (
        COMPACT_MANUAL_FIELDS if schema.layout == "compact" else TRACKING_MANUAL_FIELDS
    )

    for candidate in normalized_candidates:
        target, target_error = find_target_row(candidate, rows)
        if target_error:
            reason_text = {
                "duplicate_exact_rows": "表内存在多条完全相同的商品/样式/SKU记录",
                "existing_sku_conflict": "表内相同商品ID和样式ID已有不同商品SKU",
                "ambiguous_without_sku": "缺少商品SKU且表内相同商品ID和样式ID不唯一",
            }[target_error]
            plan["issues"].append(
                issue_for(
                    "target_conflict",
                    candidate,
                    (
                        f"商品ID {candidate['product_id']}、样式ID {candidate['style_id']}："
                        f"{reason_text}，本次未自动写入，请人工核对。"
                    ),
                )
            )
            continue

        new_row = target is None
        if new_row:
            row_number = next_row
            next_row += 1
            current_values: dict[str, object] = {}
        else:
            row_number = int(target["row"])
            current_values = target["values"]
            existing_shop = normalize_text(current_values.get("店铺名称", ""))
            if existing_shop and normalize_key(existing_shop) != normalize_key(candidate["shop_name"]):
                plan["issues"].append(
                    issue_for(
                        "target_shop_conflict",
                        candidate,
                        (
                            f"商品ID {candidate['product_id']}、样式ID {candidate['style_id']} "
                            f"在营销活动表中的店铺为“{existing_shop}”，与来源店铺"
                            f"“{candidate['shop_name']}”不同；本次未自动写入，请人工核对。"
                        ),
                    )
                )
                continue

        effective_candidate = dict(candidate)
        if not effective_candidate["sku"] and normalize_text(current_values.get("商品SKU", "")):
            effective_candidate["sku"] = normalize_text(current_values["商品SKU"])
        desired, value_issues = desired_values(
            effective_candidate,
            history,
            history_conflicts,
            costs,
            cost_conflicts,
            current_values,
        )
        plan["issues"].extend(value_issues)

        values_to_write: list[dict[str, object]] = []
        for field in AUTOMATIC_FIELDS:
            if field not in schema.columns or field not in desired:
                continue
            new_value = desired[field]
            if normalize_text(new_value) == "":
                continue
            current_value = current_values.get(field, "")
            if normalize_text(current_value) == "":
                values_to_write.append(
                    {
                        "field": field,
                        "column": schema.columns[field],
                        "value": json_value(new_value),
                        "as_text": field in {"商品ID", "样式ID", "商品SKU"},
                        "as_number": field == "产品成本" and isinstance(new_value, (int, float)),
                    }
                )
            elif not values_equal(current_value, new_value):
                plan["issues"].append(
                    issue_for(
                        "existing_value_conflict",
                        effective_candidate,
                        (
                            f"商品ID {effective_candidate['product_id']}、样式ID {effective_candidate['style_id']} 的"
                            f"{field}已有值“{normalize_text(current_value)}”，与可匹配值"
                            f"“{normalize_text(new_value)}”不同；保留原值，请人工核对。"
                        ),
                    )
                )

        missing_manual = [
            field
            for field in manual_fields
            if field in schema.columns and not normalize_text(current_values.get(field, ""))
        ]
        if missing_manual:
            plan["manual_items"].append(
                {
                    **effective_candidate,
                    "fields": missing_manual,
                    "message": (
                        f"商品ID {effective_candidate['product_id']}、样式ID {effective_candidate['style_id']} "
                        f"需人工维护：{'、'.join(missing_manual)}。"
                    ),
                }
            )

        if values_to_write:
            plan["writes"].append(
                {
                    "row": row_number,
                    "new_row": new_row,
                    "template_row": template_row,
                    "product_id": effective_candidate["product_id"],
                    "style_id": effective_candidate["style_id"],
                    "sku": effective_candidate["sku"],
                    "values": values_to_write,
                }
            )
            if new_row:
                rows.append(
                    {
                        "row": row_number,
                        "values": {
                            item["field"]: item["value"]
                            for item in values_to_write
                        },
                    }
                )
                template_row = row_number

    plan["formula_columns"] = [
        {
            "field": field,
            "column": schema.columns[field],
            "formula_r1c1": TRACKING_FORMULAS_R1C1[field],
        }
        for field in TRACKING_FORMULA_FIELDS
        if schema.layout == "tracking" and field in schema.columns
    ]
    plan["status"] = "ready"
    return plan


def apply_marketing_activity_plan(
    plan: dict[str, object],
    powershell_script: Path,
) -> dict[str, object]:
    if plan.get("status") != "ready":
        return {"status": "skipped", "changes": [], "issues": []}
    schema = plan.get("schema") or {}
    source_path = Path(str(schema.get("path", "")))
    if not source_path.exists():
        raise RuntimeError(f"营销活动表已不存在：{source_path}")
    if not plan.get("writes"):
        return {"status": "no_change", "changes": [], "issues": []}

    with tempfile.TemporaryDirectory(prefix="codex_marketing_activity_") as temp_dir_text:
        temp_dir = Path(temp_dir_text)
        staged_path = temp_dir / source_path.name
        plan_path = temp_dir / "plan.json"
        result_path = temp_dir / "result.json"
        plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
        completed = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(powershell_script),
                "-InputPath",
                str(source_path),
                "-PlanPath",
                str(plan_path),
                "-OutputPath",
                str(staged_path),
                "-ResultPath",
                str(result_path),
            ],
            check=False,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            encoding="utf-8",
            errors="replace",
        )
        if completed.returncode != 0:
            detail = completed.stderr.strip() or completed.stdout.strip()
            raise RuntimeError(f"Excel 写入营销活动表失败：{detail}")
        if not result_path.exists():
            raise RuntimeError("Excel 写入营销活动表后未返回核对结果")
        try:
            result = json.loads(result_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError as exc:
            raise RuntimeError(
                f"无法解析营销活动表写入结果：{result_path.read_text(encoding='utf-8', errors='replace')}"
            ) from exc

        backup_folder = source_path.parent / BACKUP_FOLDER_NAME
        backup_folder.mkdir(parents=True, exist_ok=True)
        timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        backup_path = backup_folder / (
            f"{source_path.stem}_营销活动更新前_{timestamp}{source_path.suffix}"
        )
        shutil.copy2(source_path, backup_path)
        try:
            os.replace(staged_path, source_path)
        except Exception:
            backup_path.unlink(missing_ok=True)
            raise
        result["backup_path"] = str(backup_path)
        result["workbook_path"] = str(source_path)
        return result


def change_details_from_plan(plan: dict[str, object]) -> list[dict[str, object]]:
    details: list[dict[str, object]] = []
    for write in plan.get("writes", []):
        details.append(
            {
                "product_id": write["product_id"],
                "style_id": write["style_id"],
                "sku": write["sku"],
                "new_row": bool(write["new_row"]),
                "fields": [item["field"] for item in write["values"]],
            }
        )
    return details
