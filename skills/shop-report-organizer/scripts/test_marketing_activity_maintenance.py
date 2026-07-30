#!/usr/bin/env python
"""Regression tests for marketing activity maintenance. Uses temporary workbooks only."""

from __future__ import annotations

import csv
import io
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

import openpyxl

import organize_reports
from marketing_activity_maintenance import (
    COMPACT_HEADERS,
    TRACKING_HEADERS,
    apply_marketing_activity_plan,
    prepare_marketing_activity_plan,
)


SCRIPT_DIR = Path(__file__).resolve().parent


def write_cost_workbook(folder: Path, rows: list[tuple[str, float]]) -> Path:
    path = folder / "成本表.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "成本"
    sheet.append(["商品编码", "规格型号", "6.11成本价", "品牌"])
    for sku, cost in rows:
        sheet.append([sku, "", cost, "成本表品牌不用于营销表"])
    workbook.save(path)
    workbook.close()
    return path


def candidate(product_id: str, style_id: str, sku: str) -> dict[str, str]:
    return {
        "shop_name": "新店铺",
        "product_id": product_id,
        "style_id": style_id,
        "sku": sku,
        "title": "测试商品",
        "source_name": "订单.csv",
    }


class MarketingActivityMaintenanceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="marketing_activity_test_")
        self.database_root = Path(self.temp.name) / "运营数据库"
        self.folder = self.database_root / "营销活动监控"
        self.folder.mkdir(parents=True)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_compact_plan_and_apply(self) -> None:
        path = self.folder / "营销活动.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "营销数据"
        sheet.append(COMPACT_HEADERS)
        sheet.append(
            [
                "STYLE-OLD",
                "900000000001",
                "历史品牌",
                "历史运营",
                "拼多多",
                "旧店铺",
                "粮油",
                "SKU-A",
                1,
                "",
                "",
                "",
            ]
        )
        workbook.save(path)
        workbook.close()
        write_cost_workbook(self.folder, [("SKU-A", 12.5)])

        plan = prepare_marketing_activity_plan(
            self.database_root,
            [candidate("977770676548", "STYLE-NEW", "SKU-A")],
        )
        self.assertEqual(plan["status"], "ready")
        self.assertEqual(len(plan["writes"]), 1)
        planned_fields = {item["field"] for item in plan["writes"][0]["values"]}
        self.assertEqual(
            planned_fields,
            {
                "店铺名称",
                "商品ID",
                "样式ID",
                "商品SKU",
                "品牌",
                "所属运营",
                "平台",
                "大类",
                "产品成本",
            },
        )

        result = apply_marketing_activity_plan(
            plan,
            SCRIPT_DIR / "update-marketing-activity.ps1",
        )
        self.assertEqual(result["status"], "updated")
        workbook = openpyxl.load_workbook(path, data_only=False)
        sheet = workbook["营销数据"]
        row = [sheet.cell(3, column).value for column in range(1, 13)]
        workbook.close()
        self.assertEqual(row[:9], [
            "STYLE-NEW",
            "977770676548",
            "历史品牌",
            "历史运营",
            "拼多多",
            "新店铺",
            "粮油",
            "SKU-A",
            12.5,
        ])
        self.assertEqual(row[9:], [None, None, None])

    def test_tracking_preserves_silent_fields_and_adds_formulas(self) -> None:
        path = self.folder / "营销活动跟踪.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "营销活动监控跟踪表"
        sheet.append(["品牌", "所属运营", "平台", "店铺名称", "活动信息", *([""] * 14), "公式计算", *([""] * 6)])
        sheet.append(["", "", "", "", *TRACKING_HEADERS[4:]])
        sheet.append(
            [
                "历史品牌",
                "历史运营",
                "拼多多",
                "旧店铺",
                "粮油",
                "900000000002",
                "STYLE-OLD",
                "SKU-B",
                8,
                "",
                "",
                "",
                "",
                "",
                20,
                "",
                18,
                "",
                "",
                '=IFERROR(M3-L3,"")',
                '=IFERROR(IF(O3+T3=0,"",O3+T3),"")',
                '=IFERROR(U3-I3,"")',
                '=IFERROR(V3/U3,"")',
                '=IFERROR(IF(Q3+T3=0,"",T3+Q3),"")',
                '=IFERROR(X3-I3,"")',
                '=IFERROR(Y3/X3,"")',
            ]
        )
        sheet.cell(4, 14).value = "KEEP-OTHER-DISCOUNT"
        sheet.cell(4, 16).value = "KEEP-MULTI-DISCOUNT"
        sheet.cell(4, 18).value = "KEEP-START"
        sheet.cell(4, 19).value = "KEEP-END"
        workbook.save(path)
        workbook.close()
        write_cost_workbook(self.folder, [("SKU-B", 10)])

        plan = prepare_marketing_activity_plan(
            self.database_root,
            [candidate("977770676549", "STYLE-NEW-B", "SKU-B")],
        )
        self.assertEqual(plan["schema"]["layout"], "tracking")
        self.assertEqual(plan["writes"][0]["row"], 4)
        result = apply_marketing_activity_plan(
            plan,
            SCRIPT_DIR / "update-marketing-activity.ps1",
        )
        formula_fields = {
            change["field"]
            for change in result["changes"]
            if change["kind"] == "formula"
        }
        self.assertEqual(
            formula_fields,
            {
                "补贴金额",
                "单件预估实收（补贴后）",
                "单件预估毛利",
                "单件预估毛利率",
                "多件预估实收（补贴后）",
                "多件预估毛利",
                "多件预估毛利率",
            },
        )

        workbook = openpyxl.load_workbook(path, data_only=False)
        sheet = workbook["营销活动监控跟踪表"]
        self.assertEqual(sheet.cell(4, 14).value, "KEEP-OTHER-DISCOUNT")
        self.assertEqual(sheet.cell(4, 16).value, "KEEP-MULTI-DISCOUNT")
        self.assertEqual(sheet.cell(4, 18).value, "KEEP-START")
        self.assertEqual(sheet.cell(4, 19).value, "KEEP-END")
        self.assertEqual(sheet.cell(4, 15).value, None)
        self.assertEqual(sheet.cell(4, 17).value, None)
        for column in range(20, 27):
            self.assertTrue(str(sheet.cell(4, column).value).startswith("="))
        workbook.close()

    def test_two_layouts_are_not_auto_selected(self) -> None:
        compact = openpyxl.Workbook()
        compact.active.append(COMPACT_HEADERS)
        compact.save(self.folder / "简版.xlsx")
        compact.close()

        tracking = openpyxl.Workbook()
        tracking.active.append(["品牌", "所属运营", "平台", "店铺名称", "活动信息", *([""] * 14), "公式计算", *([""] * 6)])
        tracking.active.append(["", "", "", "", *TRACKING_HEADERS[4:]])
        tracking.save(self.folder / "跟踪版.xlsx")
        tracking.close()

        plan = prepare_marketing_activity_plan(
            self.database_root,
            [candidate("977770676550", "STYLE-X", "SKU-X")],
        )
        self.assertEqual(plan["status"], "skipped")
        self.assertEqual(plan["writes"], [])
        self.assertTrue(any(issue["kind"] == "multiple_layouts" for issue in plan["issues"]))

    def test_organizer_dry_run_builds_mapping_and_marketing_preview(self) -> None:
        marketing_path = self.folder / "营销活动.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "营销数据"
        sheet.append(COMPACT_HEADERS)
        sheet.append(
            [
                "STYLE-HISTORY",
                "900000000003",
                "历史品牌",
                "历史运营",
                "拼多多",
                "历史店铺",
                "粮油",
                "SKU-Z",
                5,
                "",
                "",
                "",
            ]
        )
        workbook.save(marketing_path)
        workbook.close()
        write_cost_workbook(self.folder, [("SKU-Z", 7)])

        input_dir = self.database_root / "下载未分类"
        input_dir.mkdir()
        (self.database_root / "订单数据" / "测试店铺").mkdir(parents=True)
        order_path = input_dir / "订单.csv"
        with order_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(
                [
                    "订单号",
                    "订单状态",
                    "商品id",
                    "支付时间",
                    "商品",
                    "样式ID",
                    "商家编码-规格维度",
                    "店铺名称",
                ]
            )
            writer.writerow(
                [
                    "ORDER-1",
                    "已收货",
                    "988888888888",
                    "2026-07-24 08:00:00",
                    "正常商品",
                    "STYLE-Z",
                    "SKU-Z",
                    "测试店铺",
                ]
            )

        dry_run_output = io.StringIO()
        with redirect_stdout(dry_run_output):
            exit_code = organize_reports.main(
                [
                    "--database-root",
                    str(self.database_root),
                    "--input",
                    str(input_dir),
                    "--dry-run",
                ]
            )
        self.assertEqual(exit_code, 0)
        self.assertIn(
            "执行状态：仅预览，尚未正式整理；营销活动表实际新增 0 行、实际更新已有 0 行。",
            dry_run_output.getvalue(),
        )
        self.assertIn("计划新增 1 行", dry_run_output.getvalue())
        reminder_text = (self.database_root / "缺失提醒_预览.csv").read_text(
            encoding="utf-8-sig"
        )
        self.assertIn("marketing_activity_preview", reminder_text)
        self.assertIn("988888888888", reminder_text)
        with (self.database_root / "缺失提醒_预览.csv").open(
            "r", encoding="utf-8-sig", newline=""
        ) as handle:
            preview_marketing_row = next(
                row
                for row in csv.DictReader(handle)
                if row["status"] == "marketing_activity_preview"
            )
        self.assertEqual(preview_marketing_row["marketing_new_rows"], "1")
        self.assertEqual(preview_marketing_row["marketing_updated_rows"], "0")
        self.assertEqual(preview_marketing_row["marketing_unchanged_rows"], "0")
        self.assertEqual(preview_marketing_row["marketing_backup_path"], "")

        mapping = openpyxl.load_workbook(
            self.database_root / "未收录商品ID对照表_预览.xlsx",
            read_only=True,
            data_only=True,
        )
        values = list(mapping["未收录商品ID对照"].values)
        mapping.close()
        self.assertIn(
            (
                "需补导商品数据",
                "测试店铺",
                "988888888888",
                "STYLE-Z",
                "SKU-Z",
                "正常商品",
                "订单.csv",
            ),
            values,
        )

        apply_output = io.StringIO()
        with redirect_stdout(apply_output):
            apply_exit_code = organize_reports.main(
                [
                    "--database-root",
                    str(self.database_root),
                    "--input",
                    str(input_dir),
                    "--apply",
                    "--allow-missing-required",
                    "--no-update-map",
                ]
            )
        self.assertEqual(apply_exit_code, 0)
        self.assertIn("执行状态：正式整理已执行。", apply_output.getvalue())
        self.assertIn("实际新增 1 行", apply_output.getvalue())
        apply_reminder_text = (self.database_root / "缺失提醒.csv").read_text(
            encoding="utf-8-sig"
        )
        self.assertIn("marketing_activity_updated", apply_reminder_text)
        self.assertIn("实际字段明细", apply_reminder_text)
        self.assertIn("商品ID 988888888888", apply_reminder_text)
        self.assertIn("样式ID STYLE-Z", apply_reminder_text)
        self.assertIn("商品SKU SKU-Z", apply_reminder_text)
        with (self.database_root / "缺失提醒.csv").open(
            "r", encoding="utf-8-sig", newline=""
        ) as handle:
            applied_marketing_row = next(
                row
                for row in csv.DictReader(handle)
                if row["status"] == "marketing_activity_updated"
            )
        self.assertEqual(applied_marketing_row["marketing_new_rows"], "1")
        self.assertEqual(applied_marketing_row["marketing_updated_rows"], "0")
        self.assertEqual(applied_marketing_row["marketing_unchanged_rows"], "0")
        self.assertTrue(Path(applied_marketing_row["marketing_backup_path"]).exists())

        workbook = openpyxl.load_workbook(marketing_path, read_only=True, data_only=True)
        marketing_rows = list(workbook["营销数据"].values)
        workbook.close()
        self.assertEqual(marketing_rows[2][:9], (
            "STYLE-Z",
            "988888888888",
            "历史品牌",
            "历史运营",
            "拼多多",
            "测试店铺",
            "粮油",
            "SKU-Z",
            7,
        ))

    def test_chat_reminder_output_is_concise_but_audit_rows_stay_complete(self) -> None:
        product_ids = [str(900000000000 + index) for index in range(1, 9)]
        long_detail = "逐商品逐样式完整字段明细-" + ("很长" * 300)
        rows = [
            {
                "status": "new_product_id",
                "report_type": "商品数据",
                "detected_count": "8",
                "shop_name": "测试店铺",
                "source_name": "订单.csv",
                "new_product_ids": "|".join(product_ids),
                "message": long_detail,
            },
            {
                "status": "pending",
                "report_type": "",
                "detected_count": "",
                "shop_name": "",
                "source_name": "待确认.xlsx",
                "new_product_ids": "",
                "message": "文件格式未知，请确认用途。",
            },
            {
                "action": "dry-run",
                "status": "marketing_activity_preview",
                "report_type": "营销活动表",
                "detected_count": "87",
                "shop_name": "",
                "source_name": "营销活动.xlsx",
                "new_product_ids": "|".join(product_ids),
                "marketing_new_rows": "60",
                "marketing_updated_rows": "27",
                "marketing_unchanged_rows": "5",
                "marketing_backup_path": "",
                "message": long_detail,
            },
            {
                "status": "marketing_activity_manual",
                "report_type": "营销活动表",
                "detected_count": "25",
                "shop_name": "",
                "source_name": "营销活动.xlsx",
                "new_product_ids": "|".join(product_ids),
                "message": "品牌、所属运营、平台、大类、商品SKU和产品成本需要人工核对。" + long_detail,
            },
            {
                "status": "marketing_activity_manual_fields",
                "report_type": "营销活动表",
                "detected_count": "87",
                "shop_name": "",
                "source_name": "营销活动.xlsx",
                "new_product_ids": "|".join(product_ids),
                "message": "活动价、报名价、促销机制需要人工填写。" + long_detail,
            },
            {
                "status": "unrecorded_product_mapping",
                "report_type": "商品数据",
                "detected_count": "8",
                "shop_name": "",
                "source_name": "未收录商品ID对照表_预览.xlsx",
                "new_product_ids": "",
                "message": long_detail,
            },
        ]
        original_rows = [dict(row) for row in rows]
        reminder_path = self.database_root / "缺失提醒_预览.csv"

        output = io.StringIO()
        with redirect_stdout(output):
            organize_reports.print_missing_reminders(rows, reminder_path)
        chat_text = output.getvalue()

        self.assertEqual(rows, original_rows)
        self.assertNotIn(long_detail, chat_text)
        self.assertEqual(chat_text.count("[营销活动]"), 1)
        self.assertIn(
            "执行状态：仅预览，尚未正式整理；营销活动表实际新增 0 行、实际更新已有 0 行。",
            chat_text,
        )
        self.assertIn("计划新增 60 行", chat_text)
        self.assertIn("计划补充已有 27 行", chat_text)
        self.assertIn("计划无变化 5 行", chat_text)
        self.assertIn("实际新增 0 行", chat_text)
        self.assertIn("自动字段核对项 25 条", chat_text)
        self.assertIn("需人工填写记录 87 条", chat_text)
        self.assertIn("等8个", chat_text)
        self.assertIn("文件格式未知，请确认用途。", chat_text)
        self.assertIn("未收录商品ID对照表_预览.xlsx", chat_text)
        self.assertIn(f"完整提醒文件：{reminder_path}", chat_text)

    def test_marketing_summary_splits_new_updated_and_unchanged_counts(self) -> None:
        plan = {
            "status": "ready",
            "schema": {
                "path": str(self.folder / "营销活动.xlsx"),
                "layout": "compact",
            },
            "candidates": [
                candidate("900000000101", "STYLE-NEW", "SKU-NEW"),
                candidate("900000000102", "STYLE-OLD", "SKU-OLD"),
                candidate("900000000103", "STYLE-SAME", "SKU-SAME"),
            ],
            "writes": [
                {
                    "product_id": "900000000101",
                    "style_id": "STYLE-NEW",
                    "sku": "SKU-NEW",
                    "new_row": True,
                    "values": [{"field": "商品ID"}],
                },
                {
                    "product_id": "900000000102",
                    "style_id": "STYLE-OLD",
                    "sku": "SKU-OLD",
                    "new_row": False,
                    "values": [{"field": "产品成本"}],
                },
            ],
            "issues": [],
            "manual_items": [],
        }
        input_dir = self.database_root / "下载未分类"
        backup_path = self.folder / "备份" / "营销活动_更新前.xlsx"

        preview_rows = organize_reports.marketing_activity_reminder_rows(
            plan,
            "dry-run",
            input_dir,
        )
        preview = next(
            row
            for row in preview_rows
            if row["status"] == "marketing_activity_preview"
        )
        self.assertEqual(preview["marketing_new_rows"], "1")
        self.assertEqual(preview["marketing_updated_rows"], "1")
        self.assertEqual(preview["marketing_unchanged_rows"], "1")
        self.assertIn("计划新增 1 行", preview["message"])
        self.assertIn("实际新增 0 行", preview["message"])

        applied_rows = organize_reports.marketing_activity_reminder_rows(
            plan,
            "apply",
            input_dir,
            apply_result={
                "status": "updated",
                "changes": [
                    {
                        "product_id": "900000000101",
                        "style_id": "STYLE-NEW",
                        "sku": "SKU-NEW",
                        "new_row": True,
                        "field": "商品ID",
                    },
                    {
                        "product_id": "900000000102",
                        "style_id": "STYLE-OLD",
                        "sku": "SKU-OLD",
                        "new_row": False,
                        "field": "产品成本",
                    },
                ],
                "issues": [],
                "backup_path": str(backup_path),
            },
        )
        applied = next(
            row
            for row in applied_rows
            if row["status"] == "marketing_activity_updated"
        )
        self.assertEqual(applied["marketing_new_rows"], "1")
        self.assertEqual(applied["marketing_updated_rows"], "1")
        self.assertEqual(applied["marketing_unchanged_rows"], "1")
        self.assertEqual(applied["marketing_backup_path"], str(backup_path))

        output = io.StringIO()
        with redirect_stdout(output):
            organize_reports.print_missing_reminders(
                applied_rows,
                self.database_root / "缺失提醒.csv",
            )
        chat_text = output.getvalue()
        self.assertIn("执行状态：正式整理已执行。", chat_text)
        self.assertIn("实际新增 1 行", chat_text)
        self.assertIn("实际补充已有 1 行", chat_text)
        self.assertIn("未变化 1 行", chat_text)
        self.assertIn(str(backup_path), chat_text)


if __name__ == "__main__":
    unittest.main()
